#!/usr/bin/env python3
"""Generate 30-day-content-calendar.xlsx from topics.yaml (single source of truth).

Sheets:
  Calendar      — day-by-day schedule with dates, formats, hooks, status
  Topic Bank    — exploded fact rows with sources (the evidence desk)
  Long-forms    — weekly compilations and which shorts they reuse
  Production    — per-slug pipeline tracker (script/narration/render/QA/publish)

Usage:  python3 make_calendar.py [output.xlsx]
"""
import datetime, os, sys
import yaml

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX_SKILL = os.environ.get("XLSX_SKILL_DIR", "/home/z/my-project/skills/xlsx")
for sub in [XLSX_SKILL, os.path.join(XLSX_SKILL, "templates")]:
    if sub not in sys.path:
        sys.path.insert(0, sub)

from base import (FONT_NAME, HEADER_BOLD, PRIMARY, NEUTRAL_600, NEUTRAL_900,  # noqa: E402
                  ACCENT_POSITIVE, ACCENT_WARNING, setup_sheet, style_header_row,
                  style_data_row, font_caption, align_text, auto_fit_columns)
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "..", "30-day-content-calendar.xlsx")

bank = yaml.safe_load(open(os.path.join(HERE, "..", "topics.yaml")))
meta = bank["meta"]
d0 = datetime.date.fromisoformat(str(meta["start_date"]))

WEEKDAY = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]


def date_of(day):
    return d0 + datetime.timedelta(days=day - 1)


def myang_label(row):
    m = row.get("myang") or {}
    beat = m.get("beat", "")
    expr = m.get("expression", "")
    return f"{beat} ({expr})" if expr else str(beat)


def write_sheet(ws, title, headers, rows, widths=None, status_col=None, note=None):
    last_col = len(headers) + 1
    setup_sheet(ws, title=title, last_col=last_col)
    for c, h in enumerate(headers, 2):
        ws.cell(row=4, column=c, value=h)
    style_header_row(ws, row_num=4, col_start=2, col_end=last_col)
    for i, row in enumerate(rows):
        rn = 5 + i
        for c, v in enumerate(row, 2):
            ws.cell(row=rn, column=c, value=v)
        style_data_row(ws, row_num=rn, col_start=2, col_end=last_col, row_index=i)
        # status colouring
        if status_col is not None:
            cell = ws.cell(row=rn, column=status_col)
            val = str(cell.value or "")
            if val == "published":
                cell.font = Font(name=FONT_NAME, size=11, color=ACCENT_POSITIVE)
            elif val == "planned":
                cell.font = Font(name=FONT_NAME, size=11, color=NEUTRAL_600)
            elif val in ("scripted", "narrated", "rendered", "in QA"):
                cell.font = Font(name=FONT_NAME, size=11, color=ACCENT_WARNING)
    if note:
        nr = 5 + len(rows) + 2
        ws.cell(row=nr, column=2, value=note).font = font_caption()
    if widths:
        for c, w in enumerate(widths, 2):
            ws.column_dimensions[get_column_letter(c)].width = w
    else:
        auto_fit_columns(ws, min_width=8, max_width=36, header_row=4, data_start_row=5)
    ws.freeze_panes = "C5"
    return ws


wb = Workbook()

# ── Sheet 1: Calendar ────────────────────────────────────────────────────────
ws = wb.active
ws.title = "Calendar"
rows = []
for d in bank["days"]:
    dt = date_of(d["day"])
    rows.append([
        d["day"], dt.isoformat(), WEEKDAY[dt.weekday()],
        d["format"].upper(), d["title"], d["slug"], d["pillar"],
        d["hook"], myang_label(d), d.get("status", "planned"),
    ])
write_sheet(
    ws, f"{meta['channel']} — 30-Day Content Calendar (from {meta['start_date']})",
    ["Day", "Date", "Wk", "Format", "Topic", "Slug", "Pillar", "Hook (spoken line)",
     "Myang beat", "Status"],
    rows, widths=[5, 11, 5, 8, 30, 20, 10, 46, 26, 11], status_col=11,
    note=f"Niche: {meta['niche']}  |  Voice: {meta['voice']}  |  Long-form days (7/14/21/28) "
         f"replace the daily short; IG gets reel cutdowns from the long-form.")

# ── Sheet 2: Topic Bank (evidence desk) ─────────────────────────────────────
ws2 = wb.create_sheet("Topic Bank")
rows = []
for d in bank["days"]:
    for i, fw in enumerate(d.get("facts", []), 1):
        rows.append([d["day"], d["slug"], f"F{i}", fw["fact"], fw["source"]])
write_sheet(
    ws2, "Evidence Desk — every fact with its source",
    ["Day", "Slug", "#", "Fact (checked)", "Source"],
    rows, widths=[5, 20, 4, 70, 42],
    note="Rule: no fact reaches narration without a source line here. "
         "Primary papers and institutions beat aggregators.")

# ── Sheet 3: Long-forms ─────────────────────────────────────────────────────
ws3 = wb.create_sheet("Long-forms")
rows = []
for lf in bank["longform"]:
    dt = date_of(lf["day"])
    rows.append([lf["day"], dt.isoformat(), lf["slug"], lf["title"], lf["angle"],
                 ", ".join(str(x) for x in lf["reuses"]),
                 "Myang chapter cards (evidence pose)", "planned"])
write_sheet(
    ws3, "Weekly Long-form Compilations (16x9)",
    ["Day", "Date", "Slug", "Title", "Angle", "Reuses days", "Myang", "Status"],
    rows, widths=[5, 11, 24, 30, 52, 14, 30, 10], status_col=9,
    note="Assembly: expand the reused shorts' scene banks with connective narration — "
         "no new facts without new sources.")

# ── Sheet 4: Production tracker ─────────────────────────────────────────────
ws4 = wb.create_sheet("Production")
rows = []
for d in bank["days"]:
    st = d.get("status", "planned")
    if st == "published":
        done, todo = "done", "done"
    elif st == "planned":
        done, todo = "—", "todo"
    else:
        done, todo = "wip", "todo"
    rows.append([d["slug"], d["day"], "done" if st == "published" else "todo",
                 done if done == "done" else "todo",
                 done if done == "done" else "todo",
                 done if done == "done" else "todo",
                 done if done == "done" else "todo",
                 "PASS" if st == "published" else "—",
                 date_of(d["day"]).isoformat() if st == "published" else "—"])
write_sheet(
    ws4, "Production Tracker — pipeline stages per video",
    ["Slug", "Day", "Script", "Narration", "Timeline", "Render", "Export", "QA gate",
     "Publish date"],
    rows, widths=[20, 5, 9, 10, 9, 9, 9, 9, 13],
    note="Stage commands: bin/new_video.py -> polish yaml -> bin/gen_narration.py -> "
         "bin/import_narration.py -> bin/render.sh -> bin/qa_gate.py -> bin/publish_pack.py")

wb.properties.creator = "Z.ai"
os.makedirs(os.path.dirname(os.path.abspath(OUT)), exist_ok=True)
wb.save(OUT)
print(f"OK {OUT}  ({len(rows)} tracker rows, {len(bank['days'])} calendar days)")
